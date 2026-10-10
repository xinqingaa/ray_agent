"""真实解析器与 HTTP 契约；临时文件，不调用模型或业务数据库。"""
import asyncio
import hashlib
import io
import os
import threading
import zipfile
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image

from app.application.errors.exceptions import AppException, BadRequestError, ConflictError
from app.application.services.file_preview_service import FilePreviewService
from app.application.services.file_service import FileService
from app.domain.models.file import File
from app.domain.models.file_preview import PreviewQuery, TEXT_PAGE_BYTES
from app.infrastructure.external.file_preview.parser import parse_preview, read_text
from app.infrastructure.external.file_preview.worker import parse


def query(**values):
    return PreviewQuery(**values).model_dump()


def test_sales_workbook_real_values_styles_and_source_unchanged():
    path = Path(__file__).resolve().parents[4] / 'docs/plan/evidence/s01-sales-2026-10-08/outputs/销售汇总.xlsx'
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    result = parse(str(path), query())
    assert result['sheets'] == ['销售汇总', '去重后明细', '金额缺失订单']
    text = '\n'.join(cell['text'] for row in result['rows'] for cell in row)
    assert '18,640.00' in text and '31,400.00' in text
    assert result['rows'][0][0]['style']['bold']
    details = parse(str(path), query(sheet=1))
    assert details['next_row'] == 200 and len(details['rows']) == 200
    last = parse(str(path), query(sheet=1, row=200))
    assert last['next_row'] is None and len(last['rows']) == 89
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_xlsx_merged_cells_formula_and_forged_dimensions(tmp_path):
    path = tmp_path/'styled.xlsx'
    book = openpyxl.Workbook(); sheet = book.active
    sheet.append(['标题', None, None]); sheet.merge_cells('A1:C1')
    sheet['A1'].font = openpyxl.styles.Font(bold=True, color='123456')
    sheet['A1'].fill = openpyxl.styles.PatternFill('solid', fgColor='FFEEDD')
    sheet['A2'] = 0.25; sheet['A2'].number_format = '0.0%'
    sheet['A3'] = '=A2*2'; sheet.column_dimensions['A'].width=30
    book.save(path)
    result = parse(str(path), query())
    assert result['merges'] == [dict(row=0,column=0,rows=1,columns=3)]
    assert result['widths'][0] == 215
    assert result['rows'][0][0]['style']['color'] == '#123456'
    assert result['rows'][1][0]['text'] == '25.0%'
    assert result['rows'][2][0]['uncached'] and result['rows'][2][0]['text'] == '=A2*2'
    # A bogus small used-range must not erase data.
    forged=tmp_path/'forged.xlsx'
    with zipfile.ZipFile(path) as original, zipfile.ZipFile(forged,'w') as output:
        for entry in original.infolist():
            data=original.read(entry.filename)
            if entry.filename=='xl/worksheets/sheet1.xml':
                data=data.replace(b'dimension ref="A1:C3"',b'dimension ref="A1"')
            output.writestr(entry,data)
    assert parse(str(forged),query())['rows'][2][0]['text'] == '=A2*2'


def test_csv_quotes_unicode_column_and_row_paging(tmp_path):
    path=tmp_path/'quoted.csv'
    path.write_text('\ufeff名称,备注,金额\n"测试,一","两行\n内容",12.5\n'+''.join(f'行{i},正常,{i}\n' for i in range(210)))
    first=parse(str(path),query())
    assert first['rows'][1][0]['text']=='测试,一' and first['rows'][1][1]['text']=='两行\n内容'
    assert first['next_row']==200
    last=parse(str(path),query(row=200))
    assert len(last['rows'])==12 and last['next_row'] is None
    wide=tmp_path/'wide.tsv';wide.write_text('\t'.join(str(i) for i in range(70)))
    assert parse(str(wide),query())['next_column']==50
    assert parse(str(wide),query(column=50))['rows'][0][0]['text']=='50'


def test_xls_styles_and_merged_cells(tmp_path):
    import xlwt
    path=tmp_path/'legacy.xls';book=xlwt.Workbook();sheet=book.add_sheet('旧版表格')
    style=xlwt.easyxf('font: bold on; pattern: pattern solid, fore_colour light_blue;',num_format_str='#,##0.00')
    sheet.write_merge(0,0,0,2,'销售汇总',style);sheet.write(1,0,18640,style);book.save(str(path))
    result=parse(str(path),query())
    assert result['sheets']==['旧版表格'] and result['rows'][1][0]['text']=='18,640.00'
    assert result['rows'][0][0]['style']['bold'] and result['merges'][0]['columns']==3


def test_large_text_bounded_read_no_utf8_loss():
    content=(b'a'*(TEXT_PAGE_BYTES-1)+'中文'.encode()+b'x'*(6*1024*1024))
    stream=io.BytesIO(content)
    first=read_text(stream,0)
    assert first['next_offset']==TEXT_PAGE_BYTES-1 and '\ufffd' not in first['content']
    second=read_text(stream,first['next_offset'])
    assert second['content'].startswith('中文') and len(first['content'])<TEXT_PAGE_BYTES+1


def test_image_thumbnail_and_actual_subprocess(tmp_path):
    path=tmp_path/'image.png';Image.new('RGBA',(2200,1800),(0,150,180,90)).save(path)
    async def run():
        with path.open('rb') as handle:
            return await parse_preview(handle,path.name,'image-test',PreviewQuery())
    result=asyncio.run(run())
    assert result['width']==2200 and result['height']==1800
    assert result['thumbnail'].startswith('data:image/png;base64,')
    broken=tmp_path/'broken.xlsx';broken.write_bytes(b'not a workbook')
    async def invalid():
        with broken.open('rb') as handle:
            return await parse_preview(handle,broken.name,'broken-test',PreviewQuery())
    assert asyncio.run(invalid())['reason']=='文件无法解析'


def test_workbook_over_five_mb_still_previews(tmp_path):
    path=tmp_path/'large.xlsx'
    book=openpyxl.Workbook();book.active['A1']='超过旧限制仍可读';book.save(path)
    with zipfile.ZipFile(path,'a',compression=zipfile.ZIP_STORED) as archive:
        archive.writestr('padding.bin',b'x'*(6*1024*1024))
    assert path.stat().st_size>5*1024*1024
    async def run():
        with path.open('rb') as handle:
            return await parse_preview(handle,path.name,'large-workbook-test',PreviewQuery())
    assert asyncio.run(run())['rows'][0][0]['text']=='超过旧限制仍可读'


def test_cancelled_preview_waits_for_file_read_before_closing():
    started,release=threading.Event(),threading.Event()
    class SlowFile(io.BytesIO):
        def read(self,*args):
            started.set();release.wait(timeout=3)
            assert not self.closed
            return super().read(*args)
    handle=SlowFile(b'read safely')
    class Files:
        async def download_file(self,id):
            return handle,File(id=id,filename='slow.txt',size=11)
    async def run():
        service=FilePreviewService(Files(),None)
        task=asyncio.create_task(service.preview(PreviewQuery(),file_id='cancelled-test'))
        await asyncio.to_thread(started.wait,2)
        task.cancel()
        await asyncio.sleep(0)
        assert not handle.closed
        release.set()
        with pytest.raises(asyncio.CancelledError):await task
        assert handle.closed
    asyncio.run(run())


def test_project_source_revision_and_symlinks(tmp_path):
    path=tmp_path/'file.txt';path.write_text('first')
    class Projects:
        async def require_file_root(self,*args,**kwargs):return str(tmp_path)
    service=FilePreviewService(None,Projects())
    async def run():
        first=await service.preview(PreviewQuery(),identifier='p',path='file.txt',project_level=True)
        path.write_text('second')
        with pytest.raises(ConflictError):
            await service.preview(PreviewQuery(revision=first['revision']),identifier='p',path='file.txt',project_level=True)
        second=await service.preview(PreviewQuery(),identifier='p',path='file.txt',project_level=True)
        assert second['content']=='second' and second['revision']!=first['revision']
        os.symlink(path,tmp_path/'link')
        with pytest.raises((OSError,BadRequestError)):
            await service.preview(PreviewQuery(),identifier='p',path='link',project_level=True)
    asyncio.run(run())


def test_inline_range_and_sandboxed_html(tmp_path):
    from app.interfaces.endpoints.preview_routes import router
    from app.interfaces.service_dependencies import get_file_preview_service
    from app.interfaces.schemas import Response
    from fastapi.responses import JSONResponse
    class Files:
        async def download_file(self,id):
            return io.BytesIO(b'0123456789'),File(id=id,filename=id,size=10)
    service=FilePreviewService(Files(),None)
    app=FastAPI();app.include_router(router)
    app.dependency_overrides[get_file_preview_service]=lambda:service
    @app.exception_handler(AppException)
    async def error(request,exc):return JSONResponse({'msg':exc.msg},status_code=exc.status_code)
    with TestClient(app) as client:
        response=client.get('/files/test.pdf/preview/content',headers={'Range':'bytes=2-5'})
        assert response.status_code==206 and response.content==b'2345'
        assert response.headers['Content-Range']=='bytes 2-5/10'
        assert client.get('/files/test.pdf/preview/content',headers={'Range':'bytes=-3'}).content==b'789'
        assert client.get('/files/test.pdf/preview/content',headers={'Range':'bytes=20-30'}).status_code==416
        assert response.headers['Content-Security-Policy']=="sandbox; default-src 'none'"
        # 网页在不透明来源里运行脚本，不能请求本站接口或提交表单
        page=client.get('/files/report.HTML/preview/content')
        assert page.status_code==200 and page.headers['content-type']=='text/html; charset=utf-8'
        # 不支持范围读取；存储替身插在 doctype 之后
        assert page.content.startswith(b'<script>') and page.content.endswith(b'</script>0123456789')
        policy=[item.strip() for item in page.headers['Content-Security-Policy'].split(';')]
        sandbox=policy[0].split()
        assert sandbox[0]=='sandbox' and 'allow-scripts' in sandbox and 'allow-same-origin' not in sandbox
        assert "connect-src 'none'" in policy and "form-action 'none'" in policy and "frame-src 'none'" in policy
        assert page.headers['X-Content-Type-Options']=='nosniff'
        assert client.get('/files/test.js/preview/content').status_code==400
        assert client.get('/files/test.svg.txt/preview/content').status_code==400
        assert client.get('/files/test.pdf/preview?row=-1').status_code==422


def test_storage_shim_keeps_doctype_first():
    from app.interfaces.endpoints.preview_routes import STORAGE_SHIM, with_storage_shim
    assert with_storage_shim(b'\xef\xbb\xbf\n<!DOCTYPE html><html></html>') == b'\xef\xbb\xbf\n<!DOCTYPE html>'+STORAGE_SHIM+b'<html></html>'
    assert with_storage_shim(b'<html><body>x</body></html>') == STORAGE_SHIM+b'<html><body>x</body></html>'


def test_zip_duplicate_names_sanitized_and_failure_closes():
    handles=[]
    class Storage:
        async def download_file(self,id):
            stream=io.BytesIO(id.encode());handles.append(stream)
            return stream,File(id=id,filename='../相同.csv',size=len(id))
    service=FileService.__new__(FileService)
    service.file_storage=Storage()
    async def info(id):return File(id=id)
    service.get_file_info=info
    async def run():
        stream=await service.download_zip(['first','second','first'])
        try:
            with zipfile.ZipFile(stream) as archive:
                assert archive.namelist()==['相同.csv','相同 (2).csv']
                assert archive.read('相同.csv')==b'first'
                assert archive.read('相同 (2).csv')==b'second'
        finally:stream.close()
        assert all(handle.closed for handle in handles)
        async def missing(id):raise AppException(410,410,'临时图片已过期')
        service.get_file_info=missing
        with pytest.raises(AppException):await service.download_zip(['expired'])
    asyncio.run(run())


def test_download_http_head_zip_and_expired_preview():
    from app.interfaces.endpoints.file_routes import router as downloads
    from app.interfaces.endpoints.preview_routes import router as previews
    from app.interfaces.service_dependencies import get_file_service, get_file_preview_service
    from fastapi.responses import JSONResponse
    class Storage:
        async def download_file(self,id):
            return io.BytesIO(b'original bytes'),File(id=id,filename='原文件.csv',size=14)
    service=FileService.__new__(FileService);service.file_storage=Storage()
    async def info(id):
        return File(id=id,filename='原文件.csv',size=14,visual={'deleted_at':'expired'} if id=='expired' else None)
    service.get_file_info=info
    app=FastAPI();app.include_router(downloads);app.include_router(previews)
    app.dependency_overrides[get_file_service]=lambda:service
    app.dependency_overrides[get_file_preview_service]=lambda:FilePreviewService(service,None)
    @app.exception_handler(AppException)
    async def error(request,exc):return JSONResponse({'msg':exc.msg},status_code=exc.status_code)
    with TestClient(app) as client:
        assert client.head('/files/one/download').status_code==200
        assert client.get('/files/one/download').content==b'original bytes'
        assert client.head('/files/download-batch?file_id=one&file_id=two').status_code==200
        result=client.get('/files/download-batch?file_id=one&file_id=two')
        with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
            assert archive.namelist()==['原文件.csv','原文件 (2).csv']
        assert client.get('/files/expired/preview').status_code==410
        assert client.head('/files/download-batch?file_id=expired').status_code==410
