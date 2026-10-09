"""在可终止的独立进程内解析不可信文件，只输出当前阅读窗口。"""
import base64
import csv
import io
import json
import re
import sys
import zipfile
from datetime import date, datetime, time
from pathlib import Path

from app.domain.models.file_preview import (IMAGE_MAX_PIXELS, TABLE_PAGE_COLUMNS,
    TABLE_PAGE_ROWS, WORKBOOK_EXPANDED_BYTES, preview_kind)


def unavailable(reason):
    return dict(kind='unavailable', reason=reason)


def color(value, theme=None):
    if value is None:
        return None
    raw = value.rgb if value.type == 'rgb' else (theme or {}).get(value.theme) if value.type == 'theme' else None
    if not isinstance(raw, str) or not re.fullmatch(r'[0-9a-fA-F]{6,8}', raw):
        return None
    raw = raw[-6:]
    tint = value.tint or 0
    if tint:
        channels = [int(raw[i:i+2], 16) for i in (0, 2, 4)]
        raw = ''.join(f'{round(c*(1+tint) if tint < 0 else c+(255-c)*tint):02x}' for c in channels)
    return '#' + raw


def display(value, number_format='General'):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=' ') if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, (int, float)) and number_format not in ('General', '@'):
        pattern = number_format.split(';')[0]
        if re.search(r'[0#]', pattern) and not re.search(r'[Ee][+-]?[0#]|[?/]|\[[hms]', pattern):
            decimals = re.search(r'\.(0+|#+)', pattern)
            precision = min(len(decimals[1]) if decimals else 0, 12)
            percentage = '%' in pattern
            result = format(value * (100 if percentage else 1), (',' if ',' in pattern else '') + f'.{precision}f')
            if percentage:
                result += '%'
            currency = re.search(r'[$¥€£]', pattern)
            if currency:
                result = currency[0] + result
            return result
    return str(value)


def cell_style(cell, theme):
    if not cell.has_style:
        return None
    return dict(bold=bool(cell.font.b), italic=bool(cell.font.i),
        color=color(cell.font.color, theme), background=color(cell.fill.fgColor, theme) if cell.fill.patternType == 'solid' else None,
        align=cell.alignment.horizontal if cell.alignment.horizontal in ('left', 'right', 'center') else None,
        wrap=bool(cell.alignment.wrap_text),
        borders=[bool(getattr(cell.border, side) and getattr(cell.border, side).style) for side in ('top', 'right', 'bottom', 'left')])


def workbook_metadata(archive, sheet_index):
    from defusedxml.ElementTree import fromstring, iterparse
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
          'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    book = fromstring(archive.read('xl/workbook.xml'))
    sheets = book.findall('s:sheets/s:sheet', ns)
    if sheet_index >= len(sheets):
        raise ValueError('工作表不存在')
    relationships = fromstring(archive.read('xl/_rels/workbook.xml.rels'))
    targets = {item.attrib['Id']: item.attrib['Target'] for item in relationships}
    target = targets[sheets[sheet_index].attrib['{' + ns['r'] + '}id']]
    part = target.lstrip('/') if target.startswith('/') else 'xl/' + target
    from posixpath import normpath
    part = normpath(part)
    widths, merges, theme = {}, [], {}
    if 'xl/theme/theme1.xml' in archive.namelist():
        themed = fromstring(archive.read('xl/theme/theme1.xml'))
        scheme = themed.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme')
        if scheme is not None:
            # Spreadsheet theme indexes use light1/dark1/light2/dark2, unlike XML ordering.
            mapping = [1, 0, 3, 2, 4, 5, 6, 7, 8, 9, 10, 11]
            for i, position in enumerate(mapping):
                if position < len(scheme) and len(scheme[position]):
                    item = scheme[position][0]
                    theme[i] = item.attrib.get('lastClr') or item.attrib.get('val')
    with archive.open(part) as source:
        for _, node in iterparse(source, events=('end',)):
            tag = node.tag.rsplit('}', 1)[-1]
            if tag == 'col':
                start, end = int(node.attrib['min']), int(node.attrib['max'])
                for c in range(start, min(end, 16384) + 1):
                    widths[c] = min(max(float(node.attrib.get('width', 12)) * 7 + 5, 40), 600)
            elif tag == 'mergeCell' and len(merges) < 10000:
                merges.append(node.attrib['ref'])
            node.clear()
    return widths, merges, theme


def xlsx_preview(path, query):
    import openpyxl
    from openpyxl.utils.cell import range_boundaries
    with zipfile.ZipFile(path) as archive:
        parts = archive.infolist()
        if len(parts) > 10000 or sum(item.file_size for item in parts) > WORKBOOK_EXPANDED_BYTES:
            return unavailable('工作簿数据量过大')
        widths, ranges, theme = workbook_metadata(archive, query['sheet'])
    book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = book.worksheets[query['sheet']]
        values = cached.worksheets[query['sheet']]
        # reset_dimensions prevents a bogus dimension from hiding real data.
        declared_rows, declared_columns = sheet.max_row, sheet.max_column
        sheet.reset_dimensions(); values.reset_dimensions()
        start, col = query['row'], query['column']
        window = dict(min_row=start+1, max_row=start+TABLE_PAGE_ROWS+1,
                      min_col=col+1, max_col=col+TABLE_PAGE_COLUMNS)
        rows = []
        formula_missing = False
        for original, computed in zip(sheet.iter_rows(**window), values.iter_rows(**window)):
            cells = []
            for cell, result in zip(original, computed):
                formula = getattr(cell, 'data_type', '') == 'f'
                value = result.value if formula else cell.value
                missing = formula and value is None
                formula_missing |= missing
                cells.append(dict(text=display(cell.value if missing else value, getattr(cell, 'number_format', 'General')),
                    formula=str(cell.value) if formula else None, uncached=missing,
                    style=cell_style(cell, theme) if getattr(cell, 'has_style', False) else None))
            rows.append(cells)
        has_more = len(rows) > TABLE_PAGE_ROWS
        rows = rows[:TABLE_PAGE_ROWS]
        # Trim synthetic trailing rows/columns when the dimension is absent or exaggerated.
        if not has_more:
            while rows and not any(c['text'] or c['style'] for c in rows[-1]):
                rows.pop()
        last = max((i+1 for row in rows for i, c in enumerate(row) if c['text'] or c['style']), default=1)
        columns = min(TABLE_PAGE_COLUMNS, max(last, min((declared_columns or 1)-col, TABLE_PAGE_COLUMNS)))
        rows = [row[:columns] for row in rows]
        merges = []
        for ref in ranges:
            left, top, right, bottom = range_boundaries(ref)
            if left >= col+1 and top >= start+1 and right <= col+columns and bottom <= start+len(rows):
                merges.append(dict(row=top-1-start, column=left-1-col, rows=bottom-top+1, columns=right-left+1))
        return dict(kind='table', sheets=book.sheetnames, sheet=query['sheet'], row=start, column=col,
            rows=rows, widths=[widths.get(col+i+1, 120) for i in range(columns)], merges=merges,
            total_rows=declared_rows, total_columns=declared_columns,
            next_row=start+len(rows) if has_more and start+len(rows) <= 100000 else None,
            next_column=col+columns if (declared_columns or 0) > col+columns else None,
            partial=has_more or bool((declared_columns or 0) > col+columns), formula_missing=formula_missing)
    finally:
        book.close(); cached.close()


def xls_preview(path, query):
    import xlrd
    book = xlrd.open_workbook(path, on_demand=True, formatting_info=True)
    try:
        if query['sheet'] >= book.nsheets:
            return unavailable('工作表不存在')
        sheet = book.sheet_by_index(query['sheet'])
        start, col = query['row'], query['column']
        end, last = min(start+TABLE_PAGE_ROWS, sheet.nrows), min(col+TABLE_PAGE_COLUMNS, sheet.ncols)
        rows = []
        for r in range(start, end):
            row = []
            for c in range(col, last):
                cell = sheet.cell(r, c)
                xf = book.xf_list[cell.xf_index]
                font = book.font_list[xf.font_index]
                def rgb(index):
                    v = book.colour_map.get(index)
                    return '#' + ''.join(f'{ch:02x}' for ch in v) if v else None
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, book.datemode)
                row.append(dict(text=display(value, book.format_map[xf.format_key].format_str),
                    style=dict(bold=bool(font.bold), italic=bool(font.italic), color=rgb(font.colour_index),
                        background=rgb(xf.background.pattern_colour_index) if xf.background.fill_pattern == 1 else None,
                        align={1:'left',2:'center',3:'right'}.get(xf.alignment.hor_align), wrap=bool(xf.alignment.text_wrapped),
                        borders=[bool(getattr(xf.border, side+'_line_style')) for side in ('top','right','bottom','left')])))
            rows.append(row)
        merges = [dict(row=top-start, column=left-col, rows=bottom-top, columns=right-left)
            for top,bottom,left,right in sheet.merged_cells if top>=start and bottom<=end and left>=col and right<=last]
        return dict(kind='table', sheets=book.sheet_names(), sheet=query['sheet'], row=start, column=col,
            rows=rows, merges=merges, widths=[min(max(sheet.colinfo_map[c].width/256*7+5,40),600) if c in sheet.colinfo_map else 120 for c in range(col,last)],
            total_rows=sheet.nrows, total_columns=sheet.ncols,
            next_row=end if end<sheet.nrows and end<=100000 else None,
            next_column=last if last<sheet.ncols else None, partial=end<sheet.nrows or last<sheet.ncols)
    finally:
        book.release_resources()


def csv_preview(path, query, extension):
    csv.field_size_limit(65536)
    start, col = query['row'], query['column']
    rows, max_columns, more = [], 0, False
    with open(path, 'r', encoding='utf-8-sig', errors='replace', newline='') as source:
        for index, values in enumerate(csv.reader(source, delimiter='\t' if extension == '.tsv' else ',')):
            if index < start:
                continue
            if len(rows) == TABLE_PAGE_ROWS:
                more = True
                break
            max_columns = max(max_columns, len(values))
            rows.append([dict(text=value) for value in values[col:col+TABLE_PAGE_COLUMNS]])
    count = min(TABLE_PAGE_COLUMNS, max(1, max_columns-col))
    for row in rows:
        row.extend(dict(text='') for _ in range(count-len(row)))
    return dict(kind='table', sheets=[], sheet=0, row=start, column=col, rows=rows,
        widths=[160]*count, merges=[], total_rows=None, total_columns=max_columns,
        next_row=start+len(rows) if more and start+len(rows)<=100000 else None,
        next_column=col+count if max_columns>col+count else None, partial=more or max_columns>col+count)


def image_preview(path):
    from PIL import Image, ImageOps
    Image.MAX_IMAGE_PIXELS = IMAGE_MAX_PIXELS
    with Image.open(path) as original:
        if original.width * original.height > IMAGE_MAX_PIXELS:
            return unavailable('图片尺寸过大')
        width, height = original.size
        image = ImageOps.exif_transpose(original)
        image.thumbnail((1600, 1600))
        output = io.BytesIO()
        image.convert('RGBA' if 'A' in image.getbands() else 'RGB').save(output, format='PNG')
        return dict(kind='image', width=width, height=height,
            thumbnail='data:image/png;base64,'+base64.b64encode(output.getvalue()).decode())


def parse(path, query):
    extension = Path(path).suffix.lower()
    if extension == '.xlsx':
        result = xlsx_preview(path, query)
    elif extension == '.xls':
        result = xls_preview(path, query)
    elif extension in ('.csv', '.tsv'):
        result = csv_preview(path, query, extension)
    elif preview_kind(path) == 'image':
        return image_preview(path)
    else:
        return unavailable('暂不支持预览')
    budget = 1024 * 1024
    for row in result.get('rows', []):
        for cell in row:
            text = cell['text']
            visible = max(0,min(4096,budget))
            if len(text) > visible:
                cell['text'] = text[:visible]
                cell['truncated'] = True
                result['cells_truncated'] = True
            budget -= len(cell['text'])
    if query.get('input_limited'):
        result['partial'] = True
        result['input_limited'] = True
    return result


if __name__ == '__main__':
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_CPU, (10, 10))
        if sys.platform == 'linux':
            resource.setrlimit(resource.RLIMIT_AS, (768*1024*1024, 768*1024*1024))
    except (ImportError, OSError, ValueError):
        pass
    try:
        result = parse(sys.argv[1], json.loads(sys.argv[2]))
    except MemoryError:
        result = unavailable('文件数据量过大')
    except Exception:
        result = unavailable('文件无法解析')
    print(json.dumps(result, ensure_ascii=False, allow_nan=False))
